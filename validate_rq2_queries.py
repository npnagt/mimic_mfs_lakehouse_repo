#!/usr/bin/env python3
"""Validate / run the RQ2 workflow queries (docs/query_coverage/rq2_workflow_queries.sql).

Splits the SQL file on its "-- @W<nn>: <title>" markers and submits each query to Athena
against the Gold database for --dataset:

  default     EXPLAIN each query -- checks every table, column, type and function resolves
              without scanning data (cheap, and safe to run while measurements are going).
  --execute   run each query for real and record status, runtime, bytes scanned, row count
              and the first --preview-rows rows; also writes an Excel workbook with a Summary
              tab and one tab per query (W01..W30) holding its full result set (up to
              --max-excel-rows rows each). Needs openpyxl: pip install -e .[reports]

Writes docs/query_coverage/<demo|full>/rq2_queries_<explain|execute>_<UTC>.{json,md}
+ latest_<explain|execute>.{json,md} (and, with --execute, rq2_queries_execute_<UTC>.xlsx
+ latest_execute.xlsx). The console output (including any traceback) is also
saved to logs/validate_rq2_queries_<explain|execute>_<local-timestamp>.log -- no `| tee`
needed.

Note: --execute competes for Athena capacity -- don't run it while
measure_time_to_insight.py / measure_native_vs_external_join.py are measuring.

Usage:
    python validate_rq2_queries.py --dataset fulldataset                 # EXPLAIN all 30
    python validate_rq2_queries.py --dataset fulldataset --execute       # run all 30
    python validate_rq2_queries.py --dataset fulldataset --only W11,W29 --execute
"""
from __future__ import annotations

import argparse
import json
import re
import sys
import time
from datetime import datetime, timezone
from pathlib import Path

import boto3

REPO_ROOT = Path(__file__).resolve().parent
sys.path.insert(0, str(REPO_ROOT / "src"))
from mimic_lakehouse import config  # noqa: E402

SQL_FILE = REPO_ROOT / "docs" / "query_coverage" / "rq2_workflow_queries.sql"
MARKER = re.compile(r"^-- @(W\d{2}):\s*(.+)$", re.M)


class _Tee:
    """Write-through copy of a console stream into the run's log file."""

    def __init__(self, stream, log_file):
        self._stream, self._log = stream, log_file

    def write(self, text):
        self._stream.write(text)
        self._log.write(text)
        self._log.flush()
        return len(text)

    def flush(self):
        self._stream.flush()
        self._log.flush()

    def __getattr__(self, name):
        return getattr(self._stream, name)


def _start_log(mode: str) -> Path:
    """Tee stdout and stderr to logs/validate_rq2_queries_<mode>_<local-timestamp>.log."""
    logs_dir = REPO_ROOT / "logs"
    logs_dir.mkdir(parents=True, exist_ok=True)
    path = logs_dir / f"validate_rq2_queries_{mode}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.log"
    log_file = open(path, "w", encoding="utf-8")
    sys.stdout = _Tee(sys.stdout, log_file)
    sys.stderr = _Tee(sys.stderr, log_file)
    print(f"log: {path}\n$ {' '.join([Path(sys.executable).name, *sys.argv])}")
    return path


def load_queries(path: Path) -> list[dict]:
    text = path.read_text(encoding="utf-8")
    marks = list(MARKER.finditer(text))
    out = []
    for i, m in enumerate(marks):
        body = text[m.end(): marks[i + 1].start() if i + 1 < len(marks) else len(text)]
        # drop trailing section-banner comments and the terminating semicolon
        sql = body.strip().rstrip(";").strip()
        sql = re.sub(r"(\n--[^\n]*)+\s*$", "", sql).rstrip().rstrip(";")
        out.append({"id": m.group(1), "title": m.group(2).strip(), "sql": sql})
    return out


def run(ath, sql: str, database: str, output: str, workgroup: str) -> dict:
    qid = ath.start_query_execution(
        QueryString=sql, WorkGroup=workgroup,
        QueryExecutionContext={"Database": database},
        ResultConfiguration={"OutputLocation": output},
    )["QueryExecutionId"]
    while True:
        q = ath.get_query_execution(QueryExecutionId=qid)["QueryExecution"]
        state = q["Status"]["State"]
        if state in ("SUCCEEDED", "FAILED", "CANCELLED"):
            break
        time.sleep(0.5)
    stats = q.get("Statistics", {})
    return {"query_execution_id": qid, "state": state,
            "error": q["Status"].get("StateChangeReason"),
            "engine_ms": stats.get("EngineExecutionTimeInMillis"),
            "total_ms": stats.get("TotalExecutionTimeInMillis"),
            "scanned_bytes": stats.get("DataScannedInBytes")}


def fetch(ath, qid: str, keep: int) -> tuple[list[str], list[list], int]:
    """Column names, the first `keep` rows, and the total row count of a finished query."""
    cols, rows, total, token = [], [], 0, None
    while True:
        kw = {"QueryExecutionId": qid, "MaxResults": 1000}
        if token:
            kw["NextToken"] = token
        page = ath.get_query_results(**kw)
        if not cols:
            cols = [c["Name"] for c in page["ResultSet"]["ResultSetMetadata"]["ColumnInfo"]]
        data = page["ResultSet"]["Rows"]
        if total == 0 and data and [d.get("VarCharValue") for d in data[0]["Data"]] == cols:
            data = data[1:]  # header row
        for r in data:
            total += 1
            if len(rows) < keep:
                rows.append([d.get("VarCharValue") for d in r["Data"]])
        token = page.get("NextToken")
        if not token:
            return cols, rows, total


_INT = re.compile(r"-?\d{1,15}")
_FLOAT = re.compile(r"-?\d+(\.\d+)?(E[-+]?\d+)?", re.I)


def _numeric_columns(rows: list[list]) -> list[bool]:
    """Athena returns every value as a string. A column becomes numeric in Excel only if every
    non-null value in it parses as a number -- so a label column like '0', '1-2', '5+' stays text."""
    if not rows:
        return []
    return [all(v is None or _FLOAT.fullmatch(v) for v in col) for col in zip(*rows)]


def _typed(v, numeric: bool):
    if v is None or not numeric:
        return v
    return int(v) if _INT.fullmatch(v) else float(v)


def write_excel(report: dict, rows_by_id: dict[str, list[list]], path: Path) -> bool:
    """Summary sheet + one sheet per query (W01..W30) holding its result set."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        print("note: openpyxl not installed (pip install -e .[reports]) -- skipping the Excel workbook")
        return False
    hdr_fill, hdr_font = PatternFill("solid", fgColor="1F3864"), Font(bold=True, color="FFFFFF")
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws.append([f"RQ2 workflow queries -- {report['database']}"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([f"Run {report['generated_utc']} · {report['passed']}/{report['total']} succeeded · "
               f"source docs/query_coverage/rq2_workflow_queries.sql"])
    ws.append([])
    ws.append(["Id", "Workflow", "Status", "Runtime (s)", "Scanned (MB)", "Rows returned", "Rows in sheet", "Error"])
    for c in ws[4]:
        c.fill, c.font = hdr_fill, hdr_font
    for r in report["results"]:
        sheet_rows = len(rows_by_id.get(r["id"], []))
        ws.append([r["id"], r["title"], r["state"],
                   round(r["total_ms"] / 1000, 2) if r.get("total_ms") is not None else None,
                   round(r["scanned_bytes"] / 1e6, 2) if r.get("scanned_bytes") is not None else None,
                   r.get("row_count"), sheet_rows if r["state"] == "SUCCEEDED" else None,
                   None if r["state"] == "SUCCEEDED" else r.get("error")])
        if r["state"] == "SUCCEEDED":
            cell = ws.cell(row=ws.max_row, column=1)
            cell.hyperlink, cell.style = f"#'{r['id']}'!A1", "Hyperlink"
    ws.freeze_panes = "A5"
    for col, width in zip("ABCDEFGH", (7, 70, 11, 11, 12, 13, 13, 60)):
        ws.column_dimensions[col].width = width

    for r in report["results"]:
        if r["state"] != "SUCCEEDED":
            continue
        q = wb.create_sheet(r["id"])
        q.append([f"{r['id']}: {r['title']}"])
        q["A1"].font = Font(bold=True, size=12)
        shown = len(rows_by_id.get(r["id"], []))
        q.append([f"{r.get('row_count')} rows returned"
                  + (f" (first {shown} shown)" if shown < (r.get("row_count") or 0) else "")
                  + f" · {round((r.get('total_ms') or 0) / 1000, 2)} s · "
                  f"{round((r.get('scanned_bytes') or 0) / 1e6, 2)} MB scanned"])
        q.append(["← Summary"])
        q["A3"].hyperlink, q["A3"].style = "#'Summary'!A1", "Hyperlink"
        q.append([])
        q.append(r.get("columns") or [])
        for c in q[5]:
            c.fill, c.font = hdr_fill, hdr_font
            c.alignment = Alignment(wrap_text=True, vertical="top")
        numeric = _numeric_columns(rows_by_id.get(r["id"], []))
        for row in rows_by_id.get(r["id"], []):
            q.append([_typed(v, n) for v, n in zip(row, numeric)])
        q.freeze_panes = "A6"
        if r.get("columns"):
            q.auto_filter.ref = f"A5:{get_column_letter(len(r['columns']))}{max(q.max_row, 5)}"
        for i, name in enumerate(r.get("columns") or [], 1):
            longest = max([len(str(name))] + [len(str(v)) for v in (row[i - 1] for row in rows_by_id[r["id"]][:200])])
            q.column_dimensions[get_column_letter(i)].width = min(max(10, longest + 2), 60)
    wb.save(path)
    return True


def render_md(report: dict) -> str:
    L = [f"# RQ2 workflow queries -- {report['mode']} ({report['database']})", "",
         f"- **Run:** {report['generated_utc']} · mode `{report['mode']}` · "
         f"{report['passed']}/{report['total']} passed",
         f"- **Source:** `docs/query_coverage/rq2_workflow_queries.sql`", "",
         "| id | workflow | status | runtime s | scanned | rows |", "|---|---|---|---|---|---|"]
    for r in report["results"]:
        mb = f"{r['scanned_bytes'] / 1e6:.1f} MB" if r.get("scanned_bytes") is not None else ""
        rt = f"{r['total_ms'] / 1000:.1f}" if r.get("total_ms") is not None else ""
        L.append(f"| {r['id']} | {r['title']} | {r['state']} | {rt} | {mb} | {r.get('row_count', '')} |")
    for r in report["results"]:
        if r["state"] != "SUCCEEDED":
            L += ["", f"## {r['id']} -- FAILED", "", "```", str(r.get("error")), "```"]
        elif r.get("columns"):
            L += ["", f"## {r['id']} -- {r['title']}", "",
                  "| " + " | ".join(r["columns"]) + " |", "|" + "---|" * len(r["columns"])]
            for row in r["preview"]:
                L.append("| " + " | ".join("" if v is None else str(v).replace("|", "\\|").replace("\n", " ")[:120]
                                           for v in row) + " |")
    return "\n".join(L) + "\n"


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    config.add_dataset_arg(ap)
    ap.add_argument("--execute", action="store_true", help="run the queries (default: EXPLAIN only)")
    ap.add_argument("--only", default=None, help="comma-separated ids, e.g. W01,W27")
    ap.add_argument("--preview-rows", type=int, default=15, help="rows per query shown in the .md/.json report")
    ap.add_argument("--max-excel-rows", type=int, default=50000,
                    help="--execute: cap on rows written per query tab of the Excel workbook (default 50000)")
    ap.add_argument("--workgroup", default=config.ATHENA_WORKGROUP)
    ap.add_argument("--region", default=config.get("AWS_REGION"))
    args = ap.parse_args()
    mode = "execute" if args.execute else "explain"
    log_path = _start_log(mode)

    profile = config.dataset_profile(args.dataset)
    ath = boto3.Session(region_name=args.region).client("athena")
    output = f"s3://{profile.gold_bucket}/athena-results/"

    queries = load_queries(SQL_FILE)
    if args.only:
        wanted = {w.strip().upper() for w in args.only.split(",")}
        queries = [q for q in queries if q["id"] in wanted]
    print(f"{len(queries)} queries :: {mode} :: {profile.gold_database}")

    results, rows_by_id = [], {}
    for q in queries:
        sql = q["sql"] if args.execute else f"EXPLAIN {q['sql']}"
        r = {"id": q["id"], "title": q["title"], **run(ath, sql, profile.gold_database, output, args.workgroup)}
        if args.execute and r["state"] == "SUCCEEDED":
            r["columns"], kept, r["row_count"] = fetch(ath, r["query_execution_id"],
                                                       max(args.max_excel_rows, args.preview_rows))
            r["preview"] = kept[:args.preview_rows]
            rows_by_id[q["id"]] = kept[:args.max_excel_rows]
        results.append(r)
        extra = f"{r.get('row_count', '')} rows" if args.execute else ""
        secs = f"{r['total_ms'] / 1000:5.1f}s" if r.get("total_ms") is not None else "      "
        print(f"  {q['id']} {r['state']:9} {secs} {extra} {'' if r['state'] == 'SUCCEEDED' else r['error']}")

    ts = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ")
    report = {"generated_utc": ts, "mode": mode, "database": profile.gold_database,
              "total": len(results), "passed": sum(r["state"] == "SUCCEEDED" for r in results),
              "results": results}
    out = REPO_ROOT / "docs" / "query_coverage" / profile.mode
    out.mkdir(parents=True, exist_ok=True)
    for name in (f"rq2_queries_{mode}_{ts}", f"latest_{mode}"):
        (out / f"{name}.json").write_text(json.dumps(report, indent=2, default=str), encoding="utf-8")
        (out / f"{name}.md").write_text(render_md(report), encoding="utf-8")
    print(f"{report['passed']}/{report['total']} passed -> {out / f'latest_{mode}.md'}")
    if args.execute:
        xlsx = out / f"rq2_queries_execute_{ts}.xlsx"
        if write_excel(report, rows_by_id, xlsx):
            (out / "latest_execute.xlsx").write_bytes(xlsx.read_bytes())
            print(f"query results workbook: {xlsx} (+ latest_execute.xlsx)")
    print(f"log: {log_path}")
    return 0 if report["passed"] == report["total"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
